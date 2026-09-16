#!/usr/bin/env python3
"""
build_excel.py — Render a Claude-authored field-to-API mapping JSON into a
formatted Excel workbook.

Input JSON shape (--data):
{
  "module": "lending",
  "source_docx": "LOLC_Kenya_CBS_Master_Data_and_API_Requirements_v1.docx",
  "generated_at": "2026-09-08T12:00:00Z",
  "swagger_sources": [
    {"service": "lending-origination", "title": "Api Documentation", "version": "1.0", "fetched_at": "..."},
    ...
  ],
  "rows": [
    {
      "screen_id": "MOB-02",
      "mobile_screen": "Customer Details",
      "section": "Personal Details",
      "field_name": "Resident Type",
      "ui_control": "Dropdown",
      "matched_service": "lending-origination",
      "api_path": "/api/v1/customers",
      "http_method": "POST",
      "operation": "createCustomer",
      "param_location": "body",
      "request_field": "residentType",
      "response_field": "residentType",
      "data_type": "string (enum: CITIZEN, RESIDENT, NON_RESIDENT)",
      "required": "Y",
      "confidence": "High",
      "notes": "Matched by name + enum values align with dropdown options"
    },
    ...
  ]
}

Only "rows" is mandatory; any missing key on a row is rendered blank. Unknown
extra keys on a row are ignored (harmless to include for your own bookkeeping).
"confidence" should be one of: High, Medium, Low, Not Found, Blocked (case-insensitive,
substring match e.g. "High (live-verified)" still counts as High) — this drives the
row's highlight color. Anything else is left unhighlighted.

Usage:
    python build_excel.py --data mapping_rows.json --out Lending_API_Field_Mapping.xlsx
"""
import argparse
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNS = [
    ("screen_id", "Screen ID"),
    ("mobile_screen", "Mobile Screen"),
    ("section", "Section / Sub-Screen"),
    ("field_name", "Field Name"),
    ("ui_control", "UI Control"),
    ("matched_service", "Matched Service"),
    ("api_path", "API Path"),
    ("http_method", "HTTP Method"),
    ("operation", "Operation ID / Summary"),
    ("param_location", "Parameter Location"),
    ("request_field", "Request Parameter"),
    ("response_field", "Response Field"),
    ("data_type", "Data Type"),
    ("required", "Required"),
    ("confidence", "Match Confidence"),
    ("notes", "Notes"),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

CONFIDENCE_FILLS = {
    "high": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "medium": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "low": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "not found": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "blocked": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}


def resolve_fill(confidence_value):
    """Substring-match confidence text against known buckets, longest key first
    so e.g. "not found" doesn't accidentally match inside another label."""
    v = confidence_value.strip().lower()
    for key in sorted(CONFIDENCE_FILLS, key=len, reverse=True):
        if key in v:
            return CONFIDENCE_FILLS[key]
    return None

WRAP_COLS = {"notes", "field_name", "data_type", "operation", "api_path"}


def autosize(ws, columns):
    for idx, (key, header) in enumerate(columns, start=1):
        max_len = len(header)
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, min(len(str(cell.value)), 60))
        ws.column_dimensions[get_column_letter(idx)].width = max_len + 4


def build_mapping_sheet(wb, data):
    ws = wb.active
    ws.title = "Field-API Mapping"
    rows = data.get("rows", [])

    for col_idx, (key, header) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in enumerate(rows, start=2):
        fill = resolve_fill(str(row.get("confidence", "")))
        for c_idx, (key, header) in enumerate(COLUMNS, start=1):
            val = row.get(key, "")
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            wrap = key in WRAP_COLS
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=wrap)
            if fill is not None:
                cell.fill = fill

    ws.freeze_panes = "A2"
    last_col_letter = get_column_letter(len(COLUMNS))
    last_row = max(len(rows) + 1, 1)
    if len(rows) > 0:
        tbl = Table(displayName="FieldApiMapping", ref=f"A1:{last_col_letter}{last_row}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(tbl)
    autosize(ws, COLUMNS)
    return ws


def build_summary_sheet(wb, data):
    ws = wb.create_sheet("Summary")
    rows = data.get("rows", [])

    def w(r, c, v, bold=False):
        cell = ws.cell(row=r, column=c, value=v)
        if bold:
            cell.font = Font(bold=True)
        return cell

    r = 1
    w(r, 1, "Module:", bold=True); w(r, 2, data.get("module", "")); r += 1
    w(r, 1, "Source document:", bold=True); w(r, 2, data.get("source_docx", "")); r += 1
    w(r, 1, "Generated at:", bold=True); w(r, 2, data.get("generated_at", datetime.now().isoformat())); r += 1
    r += 1

    w(r, 1, "Swagger sources used", bold=True); r += 1
    w(r, 1, "Service", bold=True); w(r, 2, "Title", bold=True); w(r, 3, "Version", bold=True); w(r, 4, "Fetched at", bold=True)
    r += 1
    for src in data.get("swagger_sources", []):
        w(r, 1, src.get("service", "")); w(r, 2, src.get("title", "")); w(r, 3, src.get("version", "")); w(r, 4, src.get("fetched_at", ""))
        r += 1
    r += 1

    total = len(rows)
    counts = {}
    service_counts = {}
    for row in rows:
        conf = str(row.get("confidence", "(blank)")).strip() or "(blank)"
        counts[conf] = counts.get(conf, 0) + 1
        svc = str(row.get("matched_service", "(none)")).strip() or "(none)"
        service_counts[svc] = service_counts.get(svc, 0) + 1

    w(r, 1, "Total fields:", bold=True); w(r, 2, total); r += 1
    r += 1
    w(r, 1, "By match confidence", bold=True); r += 1
    for k, v in sorted(counts.items(), key=lambda x: -x[1]):
        w(r, 1, k); w(r, 2, v); r += 1
    r += 1
    w(r, 1, "By matched service", bold=True); r += 1
    for k, v in sorted(service_counts.items(), key=lambda x: -x[1]):
        w(r, 1, k); w(r, 2, v); r += 1

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 28
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", required=True, help="Path to mapping JSON (see module docstring for shape)")
    ap.add_argument("--out", required=True, help="Output .xlsx path")
    args = ap.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    build_mapping_sheet(wb, data)
    build_summary_sheet(wb, data)
    wb.save(args.out)

    print(f"Wrote {len(data.get('rows', []))} rows -> {args.out}")


if __name__ == "__main__":
    main()
