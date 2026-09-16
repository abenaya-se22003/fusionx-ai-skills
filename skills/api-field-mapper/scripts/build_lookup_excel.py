#!/usr/bin/env python3
"""
build_lookup_excel.py — Render a field-to-master-data-lookup mapping into a
formatted Excel workbook. Unlike build_excel.py (which documents the CREATE/
SAVE request a field's value is submitted on), this script documents the READ
side: the actual GET endpoint that supplies a dropdown/lookup's option list,
captured live from real network traffic, together with a live-captured raw
JSON response sample.

Input JSON shape (--data):
{
  "module": "lending - master data lookups (live FusionX inspect)",
  "generated_at": "2026-09-09T00:00:00Z",
  "rows": [
    {
      "screen_id": "MOB-01",
      "mobile_screen": "Lead Generation",
      "section": "Lead Details",
      "field_name": "Resident Type",
      "ui_control": "Dropdown",
      "matched_service": "comn-common",
      "api_path": "https://<host>/comn-common/residency-type/{tenant}/status/ACTIVE",
      "http_method": "GET",
      "request_parameter": "residentTypeId",
      "response_field": "<raw JSON text, as returned by the live GET call>",
      "confidence": "High",
      "notes": "..."
    },
    ...
  ]
}

Only "rows" is mandatory. "confidence" drives highlight color (substring match
against High/Medium/Low/Not Found/Blocked, case-insensitive).

Excel has a hard 32,767-character-per-cell limit. Any response_field longer
than RESPONSE_FIELD_CHAR_LIMIT is truncated with a note; the untruncated
original should be kept as a companion .json file for the user.

Usage:
    python build_lookup_excel.py --data lookup_rows.json --out Lookup_Mapping.xlsx
"""
import argparse
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RESPONSE_FIELD_CHAR_LIMIT = 30000

COLUMNS = [
    ("screen_id", "Screen ID"),
    ("mobile_screen", "Mobile Screen"),
    ("section", "Section / Sub-Screen"),
    ("field_name", "Field Name"),
    ("ui_control", "UI Control"),
    ("matched_service", "Matched Service"),
    ("api_path", "API Path"),
    ("http_method", "HTTP Method"),
    ("request_parameter", "Request Parameter"),
    ("response_field", "Response Field"),
    ("confidence", "Match Confidence"),
    ("notes", "Notes"),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

CONFIDENCE_FILLS = {
    "high": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "medium": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "low": PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid"),
    "not found": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "blocked": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    "static": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
}

WRAP_COLS = {"notes", "response_field", "api_path"}


def resolve_fill(confidence_value):
    v = confidence_value.strip().lower()
    for key in sorted(CONFIDENCE_FILLS, key=len, reverse=True):
        if key in v:
            return CONFIDENCE_FILLS[key]
    return None


def format_response_field(value):
    """Pretty-print if it parses as JSON; truncate to the Excel cell limit."""
    if value is None:
        return ""
    text = value
    if isinstance(value, (dict, list)):
        text = json.dumps(value, indent=2, ensure_ascii=False)
    elif isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith("[") or stripped.startswith("{"):
            try:
                parsed = json.loads(stripped)
                text = json.dumps(parsed, indent=2, ensure_ascii=False)
            except (ValueError, TypeError):
                text = value
        else:
            text = value
    else:
        text = str(value)

    if len(text) > RESPONSE_FIELD_CHAR_LIMIT:
        text = (
            text[:RESPONSE_FIELD_CHAR_LIMIT]
            + f"\n... [TRUNCATED at {RESPONSE_FIELD_CHAR_LIMIT} chars for Excel's cell limit "
              f"- full response is {len(text)} chars; see the companion lookup_responses/ JSON files "
              f"for the untruncated payload]"
        )
    return text


def autosize(ws, columns):
    for idx, (key, header) in enumerate(columns, start=1):
        max_len = len(header)
        cap = 80 if key not in ("response_field",) else 60
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for cell in row:
                if cell.value:
                    lines = str(cell.value).split("\n")
                    longest = max(len(line) for line in lines) if lines else 0
                    max_len = max(max_len, min(longest, cap))
        ws.column_dimensions[get_column_letter(idx)].width = max_len + 4


def build_mapping_sheet(wb, data):
    ws = wb.active
    ws.title = "Field-Lookup Mapping"
    rows = data.get("rows", [])

    for col_idx, (key, header) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in enumerate(rows, start=2):
        fill = resolve_fill(str(row.get("confidence", "")))
        for c_idx, (key, header) in enumerate(COLUMNS, start=1):
            val = row.get(key, "")
            if key == "response_field":
                val = format_response_field(val)
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            wrap = key in WRAP_COLS
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=wrap)
            if fill is not None:
                cell.fill = fill

    ws.freeze_panes = "A2"
    last_col_letter = get_column_letter(len(COLUMNS))
    last_row = max(len(rows) + 1, 1)
    if len(rows) > 0:
        tbl = Table(displayName="FieldLookupMapping", ref=f"A1:{last_col_letter}{last_row}")
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        ws.add_table(tbl)
    autosize(ws, COLUMNS)
    return ws


def build_summary_sheet(wb, data):
    ws = wb.create_sheet("Summary")
    rows = data.get("rows", [])

    def w(r, c, v, bold=False):
        cell = ws.cell(row=r, column=c, value=v)
        if bold:
            cell.font = Font(bold=True)
        return cell

    r = 1
    w(r, 1, "Module:", bold=True); w(r, 2, data.get("module", "")); r += 1
    w(r, 1, "Generated at:", bold=True); w(r, 2, data.get("generated_at", datetime.now().isoformat())); r += 1
    r += 1

    total = len(rows)
    counts = {}
    for row in rows:
        conf = str(row.get("confidence", "(blank)")).strip() or "(blank)"
        counts[conf] = counts.get(conf, 0) + 1

    w(r, 1, "Total fields:", bold=True); w(r, 2, total); r += 1
    r += 1
    w(r, 1, "By match confidence", bold=True); r += 1
    for k, v in sorted(counts.items(), key=lambda x: -x[1]):
        w(r, 1, k); w(r, 2, v); r += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50
    return ws


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", required=True, help="Path to lookup-mapping JSON (see module docstring for shape)")
    ap.add_argument("--out", required=True, help="Output .xlsx path")
    args = ap.parse_args()

    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    build_mapping_sheet(wb, data)
    build_summary_sheet(wb, data)
    wb.save(args.out)

    print(f"Wrote {len(data.get('rows', []))} rows -> {args.out}")


if __name__ == "__main__":
    main()
