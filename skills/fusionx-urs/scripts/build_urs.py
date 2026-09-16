"""Build a portable FusionX URS draft from JSON.

This provides the missing generation entry point.  It intentionally keeps content in JSON so an
agent can validate a draft before emitting Word markup rather than hiding requirements in code.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


SECTION_TITLES = [
    "Document Control", "Open Questions", "Overview/Project Description", "Flow Chart",
    "Scope", "Epic: Narrative and Statement", "Features / Stories", "Data Dictionary",
    "E2E Impact Identification Table", "Diagrams and Examples", "Annexure", "Test Scenarios",
]


def set_cell(cell, text: object, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(str(text or ""))
    run.bold = bold
    run.font.name = "Candara"
    run.font.size = Pt(9)


def table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    for cell, header in zip(result.rows[0].cells, headers):
        set_cell(cell, header, bold=True)
    for row in rows:
        cells = result.add_row().cells
        for cell, value in zip(cells, row):
            set_cell(cell, value)


def heading(document: Document, number: int, title: str) -> None:
    p = document.add_heading(level=1)
    p.style = document.styles["Heading 1"]
    run = p.add_run(f"{number}. {title}")
    run.font.name = "Candara"
    run.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)


def bullets(document: Document, values: list[str]) -> None:
    for value in values:
        p = document.add_paragraph(style="List Bullet")
        p.add_run(value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("draft_json", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    data = json.loads(args.draft_json.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Candara"
    normal.font.size = Pt(10)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover.add_run(data.get("title", "FusionX User Requirement Specification"))
    run.bold = True; run.font.name = "Candara"; run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
    doc.add_paragraph("")
    table(doc, ["Document Information", "Value"], [
        ["Client", data.get("client", "LOLC Technologies Pvt LTD")],
        ["Module", data.get("module", "")], ["Version", data.get("version", "0.1")],
        ["Drafted By", data.get("drafted_by", "")], ["Release Date", data.get("release_date", "")],
        ["Related Jira", data.get("jira", "[TO BE CONFIRMED]")],
    ])
    doc.add_page_break()

    for index, title in enumerate(SECTION_TITLES, start=1):
        heading(doc, index, title)
        key = title.lower().replace("/", "_").replace(" ", "_")
        if title == "Document Control":
            bullets(doc, data.get("assumptions", []))
        elif title == "Open Questions":
            table(doc, ["ID", "Question", "Owner", "Status"], data.get("open_questions", []))
        elif title == "Scope":
            doc.add_paragraph("In scope:"); bullets(doc, data.get("in_scope", []))
            doc.add_paragraph("Out of scope:"); bullets(doc, data.get("out_of_scope", []))
        elif title == "Features / Stories":
            for story in data.get("stories", []):
                doc.add_heading(story.get("title", "Story"), level=2)
                for label in ("user_function", "trigger", "preconditions", "action", "expected"):
                    if story.get(label):
                        doc.add_paragraph(f"{label.replace('_', ' ').title()}: {story[label]}")
        elif title == "Data Dictionary":
            table(doc, ["Feature", "Field Name", "Data Type", "Source / Retrieve From", "Constraint / Description", "Sample Data", "Data Validation", "Max Length"], data.get("data_dictionary", []))
        elif title == "Test Scenarios":
            table(doc, ["ID", "Scenario", "Expected Outcome"], data.get("test_scenarios", []))
        else:
            value = data.get(key, "")
            if value:
                doc.add_paragraph(value if isinstance(value, str) else json.dumps(value))
            else:
                doc.add_paragraph("[To be completed]")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
